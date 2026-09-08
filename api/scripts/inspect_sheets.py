import openpyxl

wb = openpyxl.load_workbook('E:/Meridian/Wajahat_Productivity_OS_Sep-Dec_2026 (1).xlsx', data_only=True)
with open('E:/Meridian/sheets_summary.txt', 'w', encoding='utf-8') as f:
    for name in wb.sheetnames:
        s = wb[name]
        f.write(f"\n=== Sheet: {name} (max_row={s.max_row}, max_column={s.max_column}) ===\n")
        for r in range(1, min(15, s.max_row + 1)):
            row_vals = [str(s.cell(r, c).value) for c in range(1, min(12, s.max_column + 1))]
            if any(v != 'None' for v in row_vals):
                f.write(f"  Row {r}: {row_vals}\n")
print("Done writing sheets_summary.txt")
